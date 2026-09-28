"""
Generator for the Apply Abroad Lab Application Tracker Pro workbook.

Run:  python generator/build_tracker.py
Output: phd-postdoc-position-tracker.xlsx in the repository root.

Author: Ali Soltanhosseini  |  applyabroadlab.com  |  databizex.com
"""
# -*- coding: utf-8 -*-
import openpyxl, datetime, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, GradientFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart3D, PieChart3D, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from sample_data import POSITIONS, TRACKING, RESPONSES
from openpyxl.worksheet.hyperlink import Hyperlink
def ilink(cell, sheet, ref):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{ref}", display=str(cell.value))


N = len(POSITIONS)
MAXROW = 200
RMAX = 119
FONT = "Calibri"

INK, NAVY, TEAL, AMBER, CRIMSON, GREEN, VIOLET, SLATE = "1F2A44","16305C","0E7C7B","D98C00","B3261E","1E7B34","5B3E90","5A6B87"
BLUE, LINE, LINK = "2E86AB","C9D3E3","0563C1"
CHART_COLS = ["16305C","0E7C7B","D98C00","B3261E","5B3E90","1E7B34","2E86AB","C2185B","7B5E2A","00838F","8E24AA","546E7A"]
TINT = {NAVY:"E3E8F2", TEAL:"DCF0EF", VIOLET:"E8E1F3", BLUE:"DDEEF5", CRIMSON:"F8E0DE", AMBER:"FBEFD6", GREEN:"DFF2E3"}

SH_DASH, SH_POS, SH_COM, SH_RES, SH_AL, SH_TPL, SH_GUIDE = ("Dashboard","Position Database","Communication Tracker",
    "Response Management","Follow-up Alerts","Email Templates","How It Works")
STATUSES = ["Sent","No Reply","Positive Reply","Neutral Reply","Rejected","Interview Invitation","Accepted"]
CLOSED   = ["Positive Reply","Accepted","Rejected"]
OUTCOMES = ["Proposal Requested","Interview Scheduled","Waitlisted","Rejected","Accepted","Pending","Closed","Info Received"]
COUNTRIES= ["Netherlands","Sweden","Denmark","Luxembourg","Germany","Switzerland","Belgium","Finland","Norway","France",
            "United Kingdom","Austria","Italy","Canada","Australia"]
PRIOS = ["High","Medium","Low"]

thin = Side(style="thin", color=LINE)
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)
POSQ, COMQ, RESQ = f"'{SH_POS}'", f"'{SH_COM}'", f"'{SH_RES}'"

wb = openpyxl.Workbook(); wb.remove(wb.active)

def newsheet(title, tab):
    ws = wb.create_sheet(title); ws.sheet_properties.tabColor = tab; ws.sheet_view.showGridLines = False; return ws

def hdr(ws, row, labels, start=1, height=34):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start+i, value=t)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    ws.row_dimensions[row].height = height

def banner(ws, lastcol, title, sub):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=lastcol)
    c = ws.cell(row=1, column=1, value=title); c.font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    c.fill = GradientFill(stop=(NAVY, TEAL)); c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=lastcol-2)
    c2 = ws.cell(row=2, column=1, value=sub); c2.font = Font(name=FONT, size=9, italic=True, color=SLATE)
    c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(start_row=2, start_column=lastcol-1, end_row=2, end_column=lastcol)
    b = ws.cell(row=2, column=lastcol-1, value="◀ Back to Dashboard"); ilink(b, SH_DASH, "A1")
    b.font = Font(name=FONT, size=9, bold=True, color=LINK, underline="single"); b.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

def widths(ws, m):
    for k, v in m.items(): ws.column_dimensions[k].width = v

def body(ws, r, c, val, align="center", wrap=False, bold=False, color=INK, size=10, fill=None):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap); cell.border = BOX
    if fill: cell.fill = PatternFill("solid", fgColor=fill)
    return cell

def linked_header(ws, ref, color):
    """Mark a header as a dashboard link destination: colored header + tinted body column."""
    c = ws[ref]; c.fill = PatternFill("solid", fgColor=color)
    col = ''.join(ch for ch in ref if ch.isalpha())
    for r in range(4, MAXROW+1):
        cc = ws[f"{col}{r}"]
        if cc.fill is None or cc.fill.fgColor is None or cc.fill.fgColor.rgb in (None, "00000000"):
            cc.fill = PatternFill("solid", fgColor=TINT[color])

def d(v): return datetime.datetime.strptime(v, "%Y-%m-%d") if v else ""

# =====================================================================
# POSITION DATABASE
# =====================================================================
ws = newsheet(SH_POS, TEAL)
banner(ws, 15, "Position Database", "One row per target position. Fill in the white cells; grey cells are calculated automatically.")
hdr(ws, 3, ["ID","Position Title","University","Country","City","Level","Research Area","Funding",
            "Deadline","Days to Deadline","Priority","Supervisor","Supervisor Email","Posting Link","Notes"])
widths(ws, {"A":6,"B":46,"C":28,"D":14,"E":15,"F":9,"G":30,"H":14,"I":12,"J":12,"K":10,"L":16,"M":24,"N":34,"O":30})
ws.freeze_panes = "A4"
for i, (pid, title, inst, ctry, city, deg, field, fund, dl, url, sname, smail) in enumerate(POSITIONS):
    r = 4+i
    body(ws, r, 1, i+1); body(ws, r, 2, title, align="left", wrap=True); body(ws, r, 3, inst, align="left")
    body(ws, r, 4, ctry); body(ws, r, 5, city); body(ws, r, 6, deg); body(ws, r, 7, field, align="left", wrap=True)
    body(ws, r, 8, fund)
    c = body(ws, r, 9, d(dl)); c.number_format = "yyyy-mm-dd"
    body(ws, r, 10, f'=IF(I{r}="","",I{r}-TODAY())', color=SLATE, fill="EEF2F8")
    body(ws, r, 11, PRIOS[i % 3]); body(ws, r, 12, sname); body(ws, r, 13, smail, align="left")
    lc = body(ws, r, 14, url, align="left"); lc.hyperlink = url; lc.font = Font(name=FONT, size=8, color=LINK, underline="single")
    body(ws, r, 15, "", align="left", wrap=True); ws.row_dimensions[r].height = 30
for r in range(4+N, MAXROW+1):
    body(ws, r, 10, f'=IF(I{r}="","",I{r}-TODAY())', color=SLATE, fill="EEF2F8")
# deadline rank helper (hidden col Q): rank of open deadlines, nearest first
ws["Q3"] = "deadline rank"; ws["Q3"].font = Font(name=FONT, size=8, color=SLATE)
for r in range(4, MAXROW+1):
    ws.cell(row=r, column=17, value=f'=IF(AND($A{r}<>"",ISNUMBER(J{r}),J{r}>=0),COUNTIFS($J$4:$J${MAXROW},">=0",$J$4:$J${MAXROW},"<"&J{r})+COUNTIF($J$4:J{r},J{r}),"")')
ws.column_dimensions["Q"].hidden = True
for dv, col in ((DataValidation(type="list", formula1='"'+",".join(PRIOS)+'"', allow_blank=True), "K"),
                (DataValidation(type="list", formula1='"'+",".join(COUNTRIES)+'"', allow_blank=True), "D"),
                (DataValidation(type="list", formula1='"PhD,Postdoc,Researcher"', allow_blank=True), "F"),
                (DataValidation(type="list", formula1='"'+"Fully Funded,Partially Funded,Unfunded,Not Stated"+'"', allow_blank=True), "H")):
    ws.add_data_validation(dv); dv.add(f"{col}4:{col}{MAXROW}")
ws.conditional_formatting.add(f"J4:J{MAXROW}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="F4D7D5"), font=Font(name=FONT, color=CRIMSON, bold=True)))
ws.conditional_formatting.add(f"J4:J{MAXROW}", CellIsRule(operator="between", formula=["0","14"], fill=PatternFill("solid", fgColor="FBEBD0"), font=Font(name=FONT, color=AMBER, bold=True)))
ws.conditional_formatting.add(f"J4:J{MAXROW}", CellIsRule(operator="greaterThan", formula=["14"], fill=PatternFill("solid", fgColor="DDF0E1"), font=Font(name=FONT, color=GREEN)))
# whole-row highlight for deadlines within 30 days (the dashboard's "Upcoming Deadlines" links land here)
ws.conditional_formatting.add(f"A4:O{MAXROW}", FormulaRule(formula=['AND(ISNUMBER($J4),$J4>=0,$J4<=30)'], fill=PatternFill("solid", fgColor="FFF4D6")))

# =====================================================================
# COMMUNICATION TRACKER
# =====================================================================
wc = newsheet(SH_COM, NAVY)
banner(wc, 15, "Communication Tracker", "Pick the position ID from the drop-down; title, university and country fill in automatically. Dates as yyyy-mm-dd. The follow-up clock stops when a case is closed.")
hdr(wc, 3, ["ID","Position Title","University","Country","First Email","Template","CV Version","Proposal Version",
            "Follow-up 1","Follow-up 2","Follow-up 3","Last Contact","Days Since","Next Follow-up","Status"])
widths(wc, {"A":6,"B":44,"C":26,"D":14,"E":13,"F":12,"G":11,"H":13,"I":13,"J":13,"K":13,"L":13,"M":11,"N":14,"O":19})
wc.freeze_panes = "A4"
for r in range(4, MAXROW+1):
    i = r-4
    if i < N:
        e1, tpl, cv, pr, f1, f2, f3, st = TRACKING[i]; body(wc, r, 1, i+1)
    for col, src in ((2,"B"),(3,"C"),(4,"D")):
        body(wc, r, col, f'=IFERROR(INDEX({POSQ}!{src}:{src},MATCH($A{r},{POSQ}!$A:$A,0)),"")',
             align="left" if col != 4 else "center", wrap=(col == 2), color=SLATE, fill="EEF2F8")
    if i < N:
        for col, val in ((5,e1),(9,f1),(10,f2),(11,f3)):
            c = body(wc, r, col, d(val)); c.number_format = "yyyy-mm-dd"
        body(wc, r, 6, tpl); body(wc, r, 7, cv); body(wc, r, 8, pr); body(wc, r, 15, st)
        wc.row_dimensions[r].height = 28
    c = body(wc, r, 12, f'=IF(COUNT(E{r},I{r},J{r},K{r})=0,"",MAX(E{r},I{r},J{r},K{r}))', color=SLATE, fill="EEF2F8"); c.number_format = "yyyy-mm-dd"
    closed = ",".join([f'O{r}="{s}"' for s in CLOSED])
    body(wc, r, 13, f'=IF(OR($A{r}="",NOT(ISNUMBER(L{r}))),"",IF(OR({closed}),"",TODAY()-L{r}))', color=SLATE, fill="EEF2F8")
    c = body(wc, r, 14, f'=IF(OR($A{r}="",NOT(ISNUMBER(L{r}))),"",IF(OR({closed}),"",L{r}+14))', color=SLATE, fill="EEF2F8"); c.number_format = "yyyy-mm-dd"
    wc.cell(row=r, column=17, value=f'=IF(AND($A{r}<>"",ISNUMBER(M{r})),COUNTIF($M$4:$M${MAXROW},">"&M{r})+COUNTIF($M$4:M{r},M{r}),"")')
wc["Q3"] = "urgency rank"; wc["Q3"].font = Font(name=FONT, size=8, color=SLATE); wc.column_dimensions["Q"].hidden = True
for dv, col in ((DataValidation(type="list", formula1='"'+",".join(STATUSES)+'"', allow_blank=True), "O"),
                (DataValidation(type="list", formula1='"'+"Template A,Template B,Template C,Template D,Template E,Custom"+'"', allow_blank=True), "F"),
                (DataValidation(type="list", formula1=f'={POSQ}!$A$4:$A${MAXROW}', allow_blank=True), "A")):
    wc.add_data_validation(dv); dv.add(f"{col}4:{col}{MAXROW}")
wc.conditional_formatting.add(f"M4:M{MAXROW}", CellIsRule(operator="greaterThan", formula=["14"], fill=PatternFill("solid", fgColor="F4D7D5"), font=Font(name=FONT, color=CRIMSON, bold=True)))
wc.conditional_formatting.add(f"M4:M{MAXROW}", CellIsRule(operator="between", formula=["7","14"], fill=PatternFill("solid", fgColor="FBEBD0"), font=Font(name=FONT, color=AMBER, bold=True)))
# whole-row highlight for overdue follow-ups (dashboard "Needs Follow-up Now" links land here)
wc.conditional_formatting.add(f"A4:O{MAXROW}", FormulaRule(formula=['AND(ISNUMBER($M4),$M4>14)'], fill=PatternFill("solid", fgColor="FBE4E2")))

# =====================================================================
# RESPONSE MANAGEMENT
# =====================================================================
wr = newsheet(SH_RES, VIOLET)
banner(wr, 9, "Response Management", "Log every reply you receive. Pick the position ID from the drop-down. Update the Status in the Communication Tracker as well so the follow-up clock stops.")
hdr(wr, 3, ["ID","Position Title","University","Reply Date","Reply Status","Reply Summary","Meeting Date","Outcome","Next Action"])
widths(wr, {"A":6,"B":42,"C":26,"D":13,"E":19,"F":48,"G":13,"H":20,"I":36})
wr.freeze_panes = "A4"
for r in range(4, RMAX+1):
    i = r-4
    for col, src in ((2,"B"),(3,"C")):
        body(wr, r, col, f'=IFERROR(INDEX({POSQ}!{src}:{src},MATCH($A{r},{POSQ}!$A:$A,0)),"")', align="left", wrap=(col == 2), color=SLATE, fill="EEF2F8")
    if i < len(RESPONSES):
        idx, rdate, rstat, summ, mdate, outc, nxt = RESPONSES[i]
        body(wr, r, 1, idx+1)
        c = body(wr, r, 4, d(rdate)); c.number_format = "yyyy-mm-dd"
        body(wr, r, 5, rstat); body(wr, r, 6, summ, align="left", wrap=True)
        c = body(wr, r, 7, d(mdate)); c.number_format = "yyyy-mm-dd"
        body(wr, r, 8, outc); body(wr, r, 9, nxt, align="left", wrap=True); wr.row_dimensions[r].height = 30
for dv, col in ((DataValidation(type="list", formula1='"'+",".join(STATUSES)+'"', allow_blank=True), "E"),
                (DataValidation(type="list", formula1='"'+",".join(OUTCOMES)+'"', allow_blank=True), "H"),
                (DataValidation(type="list", formula1=f'={POSQ}!$A$4:$A${MAXROW}', allow_blank=True), "A")):
    wr.add_data_validation(dv); dv.add(f"{col}4:{col}{RMAX}")
wr.conditional_formatting.add(f"A4:I{RMAX}", FormulaRule(formula=['$E4="Interview Invitation"'], fill=PatternFill("solid", fgColor="FBEFD6")))
wr.conditional_formatting.add(f"A4:I{RMAX}", FormulaRule(formula=['$E4="Accepted"'], fill=PatternFill("solid", fgColor="DFF2E3")))

# =====================================================================
# FOLLOW-UP ALERTS
# =====================================================================
wa = newsheet(SH_AL, AMBER)
banner(wa, 8, "Follow-up Alerts", "Red = more than 14 days, send a follow-up today   |   Amber = 7 to 14 days, prepare it   |   Green = under 7 days, keep watching")
hdr(wa, 3, ["#","Position Title","University","Days Since","Last Contact","Next Follow-up","Status","Recommended Action"])
widths(wa, {"A":6,"B":42,"C":26,"D":11,"E":13,"F":14,"G":19,"H":46})
wa.freeze_panes = "A4"
for r in range(4, MAXROW+1):
    s = r
    body(wa, r, 1, f'=IF({COMQ}!B{s}="","",{r-3})', color=SLATE)
    body(wa, r, 2, f'=IF({COMQ}!B{s}="","",{COMQ}!B{s})', align="left", wrap=True)
    body(wa, r, 3, f'=IF({COMQ}!C{s}="","",{COMQ}!C{s})', align="left")
    body(wa, r, 4, f'=IFERROR(IF(ISNUMBER({COMQ}!M{s}),{COMQ}!M{s},""),"")', bold=True)
    c = body(wa, r, 5, f'=IFERROR({COMQ}!L{s},"")'); c.number_format = "yyyy-mm-dd"
    c = body(wa, r, 6, f'=IFERROR(IF(ISNUMBER({COMQ}!N{s}),{COMQ}!N{s},""),"")'); c.number_format = "yyyy-mm-dd"
    body(wa, r, 7, f'=IF({COMQ}!O{s}="","",{COMQ}!O{s})')
    closed = ",".join([f'{COMQ}!O{s}="{x}"' for x in CLOSED])
    body(wa, r, 8, f'=IF({COMQ}!B{s}="","",IF(OR({closed}),"Case closed - no action needed",'
                   f'IF(ISNUMBER({COMQ}!M{s}),IF({COMQ}!M{s}>14,"URGENT - send a follow-up today",'
                   f'IF({COMQ}!M{s}>=7,"Prepare follow-up - send within "&(14-{COMQ}!M{s})&" days",'
                   f'"Monitoring - next follow-up in "&(14-{COMQ}!M{s})&" days")),"")))', align="left", wrap=True)
    wa.row_dimensions[r].height = 24
wa.conditional_formatting.add(f"A4:H{MAXROW}", FormulaRule(formula=['AND(ISNUMBER($D4),$D4>14)'], fill=PatternFill("solid", fgColor="FBE4E2")))
wa.conditional_formatting.add(f"A4:H{MAXROW}", FormulaRule(formula=['AND(ISNUMBER($D4),$D4>=7,$D4<=14)'], fill=PatternFill("solid", fgColor="FDF3E0")))
wa.conditional_formatting.add(f"A4:H{MAXROW}", FormulaRule(formula=['AND(ISNUMBER($D4),$D4<7)'], fill=PatternFill("solid", fgColor="E8F5EC")))

# =====================================================================
# DASHBOARD
# =====================================================================
wd = wb.create_sheet(SH_DASH, 0); wd.sheet_view.showGridLines = False; wd.sheet_properties.tabColor = CRIMSON
wd.merge_cells("A1:N1"); c = wd["A1"]; c.value = "ApplyAbroadLab.com  |  PhD & Postdoc Application Tracker"
c.font = Font(name=FONT, size=16, bold=True, color="FFFFFF", underline="single"); c.fill = GradientFill(stop=(NAVY, TEAL))
c.alignment = Alignment(horizontal="center", vertical="center"); c.hyperlink = "https://applyabroadlab.com"; wd.row_dimensions[1].height = 42
wd.merge_cells("A2:N2"); c = wd["A2"]; c.value = "Command center for a graduate application campaign  |  Every card, chart and list below recalculates automatically  |  Click any card or row to jump to its source"
c.font = Font(name=FONT, size=9, italic=True, color=SLATE); c.alignment = Alignment(horizontal="center", vertical="center"); wd.row_dimensions[2].height = 20
widths(wd, {"A":4,"B":17,"C":17,"D":17,"E":17,"F":17,"G":17,"H":17,"I":3,"J":30,"K":12,"L":3,"M":3,"N":3})

def section(ws, ref, text, color=NAVY, link=None):
    c = ws[ref]; c.value = text; c.font = Font(name=FONT, size=11, bold=True, color=color if not link else LINK, underline="single" if link else None)
    c.alignment = Alignment(horizontal="left", vertical="center")
    if link: ilink(c, *link)
    return c

# ---- KPI cards (each label links to its source column; that column is colour-marked there) ----
section(wd, "A4", "KEY METRICS")
KPI = [
    ("Target\nPositions",   f'=COUNTA({POSQ}!$B$4:$B${MAXROW})',                              NAVY,    ws, "B3"),
    ("Emails\nSent",        f'=COUNTA({COMQ}!$E$4:$E${MAXROW})',                              TEAL,    wc, "E3"),
    ("Replies\nReceived",   f'=COUNTA({RESQ}!$D$4:$D${RMAX})',                                VIOLET,  wr, "D3"),
    ("Response\nRate",      f'=IFERROR(COUNTA({RESQ}!$D$4:$D${RMAX})/COUNTA({COMQ}!$E$4:$E${MAXROW}),0)', BLUE, wr, "E3"),
    ("Overdue\nFollow-ups", f'=COUNTIF({COMQ}!$M$4:$M${MAXROW},">14")',                       CRIMSON, wc, "M3"),
    ("Upcoming\nInterviews",f'=COUNTIF({RESQ}!$E$4:$E${RMAX},"Interview Invitation")',        AMBER,   wr, "G3"),
    ("Final\nAcceptances",  f'=COUNTIF({RESQ}!$E$4:$E${RMAX},"Accepted")',                    GREEN,   wr, "H3"),
]
for i, (label, formula, color, tgt_ws, tgt_ref) in enumerate(KPI):
    col = 2+i
    lc = wd.cell(row=5, column=col, value=label); lc.font = Font(name=FONT, size=9, bold=True, color="FFFFFF", underline="single")
    lc.fill = PatternFill("solid", fgColor=color); lc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); lc.border = BOX
    ilink(lc, tgt_ws.title, tgt_ref)
    vc = wd.cell(row=6, column=col, value=formula); vc.font = Font(name=FONT, size=22, bold=True, color=color)
    vc.alignment = Alignment(horizontal="center", vertical="center"); vc.fill = GradientFill(stop=("FFFFFF", TINT[color])); vc.border = BOX
    ilink(vc, tgt_ws.title, tgt_ref)
    if "Rate" in label: vc.number_format = "0%"
    linked_header(tgt_ws, tgt_ref, color)
wd.row_dimensions[5].height = 34; wd.row_dimensions[6].height = 48

# ---- Quick navigation (J4:K11) ----
section(wd, "J4", "QUICK NAVIGATION")
NAV = [(SH_POS,"A3",TEAL),(SH_COM,"A3",NAVY),(SH_RES,"A3",VIOLET),(SH_AL,"A3",AMBER),(SH_TPL,"A3",GREEN),(SH_GUIDE,"A3",SLATE)]
for i, (name, ref, color) in enumerate(NAV):
    r = 5+i; wd.merge_cells(start_row=r, start_column=10, end_row=r, end_column=11)
    c = wd.cell(row=r, column=10, value=f"▶  {name}"); ilink(c, name, ref)
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF", underline="single"); c.fill = PatternFill("solid", fgColor=color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1); c.border = BOX; wd.row_dimensions[r].height = 22

# ---- hidden chart data (cols P:Q) ----
wd["P1"] = "Chart data (auto)"; wd["P1"].font = Font(name=FONT, size=9, bold=True, color=SLATE)
wd["P3"], wd["Q3"] = "Country", "Count"
COUNTRY_LIST = ["Netherlands","Sweden","Denmark","Luxembourg","Switzerland","Germany","Belgium","Finland","Norway","France","United Kingdom"]
for i, ctry in enumerate(COUNTRY_LIST):
    wd.cell(row=4+i, column=16, value=ctry); wd.cell(row=4+i, column=17, value=f'=COUNTIF({POSQ}!$D$4:$D${MAXROW},$P{4+i})')
C_END = 3+len(COUNTRY_LIST)
S_START = C_END+3; wd.cell(row=S_START, column=16, value="Status"); wd.cell(row=S_START, column=17, value="Count")
for i, s in enumerate(STATUSES):
    wd.cell(row=S_START+1+i, column=16, value=s); wd.cell(row=S_START+1+i, column=17, value=f'=COUNTIF({COMQ}!$O$4:$O${MAXROW},$P{S_START+1+i})')
S_END = S_START+len(STATUSES)
D_START = S_END+3; wd.cell(row=D_START, column=16, value="Level"); wd.cell(row=D_START, column=17, value="Count")
for i, dg in enumerate(["PhD","Postdoc"]):
    wd.cell(row=D_START+1+i, column=16, value=dg); wd.cell(row=D_START+1+i, column=17, value=f'=COUNTIF({POSQ}!$F$4:$F${MAXROW},$P{D_START+1+i})')
D_END = D_START+2
F_START = D_END+3; wd.cell(row=F_START, column=16, value="Stage"); wd.cell(row=F_START, column=17, value="Count")
FUNNEL = [("1. Emails sent", f'=COUNTA({COMQ}!$E$4:$E${MAXROW})'), ("2. Replies received", f'=COUNTA({RESQ}!$D$4:$D${RMAX})'),
          ("3. Interview invitations", f'=COUNTIF({RESQ}!$E$4:$E${RMAX},"Interview Invitation")'), ("4. Final acceptances", f'=COUNTIF({RESQ}!$E$4:$E${RMAX},"Accepted")')]
for i, (lab, f) in enumerate(FUNNEL):
    wd.cell(row=F_START+1+i, column=16, value=lab); wd.cell(row=F_START+1+i, column=17, value=f)
F_END = F_START+len(FUNNEL)
for r in range(1, F_END+2):
    for col in (16,17):
        if wd.cell(row=r, column=col).value is not None: wd.cell(row=r, column=col).font = Font(name=FONT, size=9, color=SLATE)
for col in "PQ": wd.column_dimensions[col].hidden = True

# ---- charts (3D, per-point colours) ----
def labels(ch, pct=False):
    ch.dLbls = DataLabelList(); ch.dLbls.showVal = True; ch.dLbls.showPercent = pct
    ch.dLbls.showCatName = False; ch.dLbls.showSerName = False; ch.dLbls.showLegendKey = False; ch.dLbls.showLeaderLines = pct

def style_series(ser, colors):
    for idx, col in enumerate(colors):
        dp = DataPoint(idx=idx); dp.graphicalProperties = GraphicalProperties(solidFill=col); dp.graphicalProperties.line.solidFill = "FFFFFF"; ser.data_points.append(dp)

section(wd, "B8", "POSITIONS BY COUNTRY  →  Position Database", link=(SH_POS,"D3")); linked_header(ws, "D3", BLUE)
section(wd, "F8", "OUTREACH STATUS  →  Communication Tracker", link=(SH_COM,"O3")); linked_header(wc, "O3", VIOLET)
ch1 = BarChart3D(); ch1.type = "col"; ch1.style = 10; ch1.title = "Positions by Country"; ch1.y_axis.title = "Positions"
ch1.height, ch1.width = 8.5, 12.4; ch1.gapDepth = 60; ch1.gapWidth = 60
ch1.add_data(Reference(wd, min_col=17, min_row=3, max_row=C_END), titles_from_data=True); ch1.set_categories(Reference(wd, min_col=16, min_row=4, max_row=C_END))
style_series(ch1.series[0], CHART_COLS[:len(COUNTRY_LIST)]); labels(ch1); ch1.legend = None
wd.add_chart(ch1, "B9")
ch2 = PieChart3D(); ch2.title = "Outreach Status"; ch2.height, ch2.width = 8.5, 9.4
ch2.add_data(Reference(wd, min_col=17, min_row=S_START, max_row=S_END), titles_from_data=True); ch2.set_categories(Reference(wd, min_col=16, min_row=S_START+1, max_row=S_END))
style_series(ch2.series[0], [TEAL, SLATE, GREEN, AMBER, CRIMSON, VIOLET, "2E86AB"]); labels(ch2, True)
wd.add_chart(ch2, "F9")

section(wd, "B27", "APPLICATION FUNNEL  →  Response Management", link=(SH_RES,"E3"))
section(wd, "F27", "PHD VS POSTDOC  →  Position Database", link=(SH_POS,"F3")); linked_header(ws, "F3", AMBER)
ch3 = BarChart3D(); ch3.type = "bar"; ch3.style = 10; ch3.title = "Application Funnel"; ch3.height, ch3.width = 8.0, 12.4; ch3.gapDepth = 80; ch3.gapWidth = 50
ch3.add_data(Reference(wd, min_col=17, min_row=F_START, max_row=F_END), titles_from_data=True); ch3.set_categories(Reference(wd, min_col=16, min_row=F_START+1, max_row=F_END))
style_series(ch3.series[0], [NAVY, TEAL, AMBER, GREEN]); labels(ch3); ch3.legend = None
wd.add_chart(ch3, "B28")
ch4 = PieChart3D(); ch4.title = "PhD vs Postdoc"; ch4.height, ch4.width = 8.0, 9.4
ch4.add_data(Reference(wd, min_col=17, min_row=D_START, max_row=D_END), titles_from_data=True); ch4.set_categories(Reference(wd, min_col=16, min_row=D_START+1, max_row=D_END))
style_series(ch4.series[0], [NAVY, AMBER]); labels(ch4, True)
wd.add_chart(ch4, "F28")

# ---- Upcoming deadlines (J13:K23) – top 8 open deadlines, each row links to its Position Database row ----
section(wd, "J13", "UPCOMING DEADLINES", color=AMBER)
for i, t in enumerate(["University  |  Position","Days"]):
    c = wd.cell(row=14, column=10+i, value=t); c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=AMBER)
    c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BOX
for k in range(1, 9):
    r = 14+k
    rowf = f'MATCH({k},{POSQ}!$Q$4:$Q${MAXROW},0)'
    c = wd.cell(row=r, column=10, value=f'=IFERROR(HYPERLINK("#\'{SH_POS}\'!A"&({rowf}+3),INDEX({POSQ}!$C$4:$C${MAXROW},{rowf})&"  |  "&INDEX({POSQ}!$B$4:$B${MAXROW},{rowf})),"")')
    c.font = Font(name=FONT, size=8, color=LINK, underline="single"); c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BOX
    c = wd.cell(row=r, column=11, value=f'=IFERROR(INDEX({POSQ}!$J$4:$J${MAXROW},{rowf}),"")')
    c.font = Font(name=FONT, size=10, bold=True, color=CRIMSON); c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BOX
    wd.row_dimensions[r].height = 30
wd.conditional_formatting.add("K15:K22", CellIsRule(operator="lessThanOrEqual", formula=["14"], fill=PatternFill("solid", fgColor="FBE4E2")))
wd.conditional_formatting.add("K15:K22", CellIsRule(operator="greaterThan", formula=["14"], fill=PatternFill("solid", fgColor="E8F5EC"), font=Font(name=FONT, bold=True, color=GREEN)))

# ---- Needs follow-up now (A46:G54) – each row links to its Communication Tracker row ----
UR = 46
section(wd, f"A{UR}", "NEEDS FOLLOW-UP NOW  —  top 8 most overdue, pulled live from the Communication Tracker", color=CRIMSON)
for i, t in enumerate(["#","Position Title","University","Country","Days Since","Status","Next Follow-up"]):
    c = wd.cell(row=UR+1, column=1+i, value=t); c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=CRIMSON)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
wd.row_dimensions[UR+1].height = 26
for k in range(1, 9):
    r = UR+1+k
    rowf = f'MATCH({k},{COMQ}!$Q$4:$Q${MAXROW},0)'
    c = wd.cell(row=r, column=1, value=k); c.font = Font(name=FONT, size=9, color=SLATE); c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BOX
    c = wd.cell(row=r, column=2, value=f'=IFERROR(HYPERLINK("#\'{SH_COM}\'!A"&({rowf}+3),INDEX({COMQ}!$B$4:$B${MAXROW},{rowf})),"")')
    c.font = Font(name=FONT, size=9, color=LINK, underline="single"); c.alignment = Alignment(horizontal="left", vertical="center"); c.border = BOX
    for col, src in ((3,"C"),(4,"D"),(5,"M"),(6,"O"),(7,"N")):
        c = wd.cell(row=r, column=col, value=f'=IFERROR(INDEX({COMQ}!${src}$4:${src}${MAXROW},{rowf}),"")')
        c.font = Font(name=FONT, size=9, color=INK, bold=(src == "M")); c.alignment = Alignment(horizontal="left" if col == 3 else "center", vertical="center"); c.border = BOX
        if src == "N": c.number_format = "yyyy-mm-dd"
    wd.row_dimensions[r].height = 22
wd.conditional_formatting.add(f"E{UR+2}:E{UR+9}", CellIsRule(operator="greaterThan", formula=["14"], fill=PatternFill("solid", fgColor="FBE4E2"), font=Font(name=FONT, bold=True, color=CRIMSON)))
wd.conditional_formatting.add(f"E{UR+2}:E{UR+9}", CellIsRule(operator="between", formula=["7","14"], fill=PatternFill("solid", fgColor="FDF3E0"), font=Font(name=FONT, bold=True, color=AMBER)))
wd.merge_cells(start_row=UR+10, start_column=1, end_row=UR+10, end_column=8)
nc = wd.cell(row=UR+10, column=1, value="Tip: update Status and dates in the Communication Tracker and deadlines in the Position Database - every card, chart and list on this page recalculates automatically. The full alert list is in Follow-up Alerts.")
nc.font = Font(name=FONT, size=8, italic=True, color=SLATE); nc.alignment = Alignment(horizontal="left", vertical="center")

# =====================================================================
# EMAIL TEMPLATES
# =====================================================================
wt = newsheet(SH_TPL, GREEN)
banner(wt, 5, "Email Templates", "Replace every [bracketed value] with your own details before sending. Keep the first email under 200 words.")
hdr(wt, 3, ["Code","Template","When to Use","Subject Line","Email Body"])
widths(wt, {"A":11,"B":24,"C":30,"D":42,"E":92}); wt.freeze_panes = "A4"
TPL = [
 ("Template A","Cold Email - First Contact","First email to a new professor",
  "PhD/Postdoc Inquiry - [Research Area] - [Your Name]",
  "Dear Professor [Last Name],\n\nMy name is [Your Name]. I hold a [MSc/PhD] in [Field] from [Your University] and have been following your group's work on [specific topic] with great interest.\n\nYour recent paper \"[Paper Title]\" ([Year]) caught my attention, particularly [one specific point from the paper]. It connects closely with my own experience in [your relevant experience].\n\nIf there is an opening for a [PhD/postdoc] position in your group, I would be glad to send my CV and a short research proposal for your consideration.\n\nKind regards,\n[Full Name]\n[Email] | [LinkedIn]"),
 ("Template B","First Follow-up","14 days after the first email, no reply",
  "Following Up - PhD/Postdoc Inquiry - [Your Name]",
  "Dear Professor [Last Name],\n\nI hope this message finds you well. I am following up on my email of [date of first email].\n\nI understand your schedule is very full. I simply wanted to ask whether there might be any opportunity to join your group as a [PhD student/postdoc].\n\nI would be happy to provide any further documents you may need.\n\nKind regards,\n[Full Name]"),
 ("Template C","Second and Final Follow-up","14 days after the first follow-up",
  "Final Follow-up - [Your Name] - [Research Area]",
  "Dear Professor [Last Name],\n\nApologies for writing once more; this will be my last follow-up.\n\nIf there are no openings in your group at the moment, I would be grateful if you could point me toward a colleague who may be recruiting in a related area.\n\nThank you for your time.\n\nKind regards,\n[Full Name]"),
 ("Template D","Reply to Expressed Interest","The professor replied positively",
  "Re: [Original Subject] - Research Proposal & Materials",
  "Dear Professor [Last Name],\n\nThank you very much for your kind reply. I am delighted that my research interests align with the goals of your group.\n\nPlease find attached:\n- Full CV ([version])\n- Two-page research proposal\n- [Other document]\n\nI would be glad to arrange an online conversation at any time that suits you.\n\nKind regards,\n[Full Name]"),
 ("Template E","Thank-you After Interview","Within 24 hours of the meeting",
  "Thank You - Our Meeting on [Date] - [Your Name]",
  "Dear Professor [Last Name],\n\nThank you for taking the time to meet with me today. I found our discussion about [specific topic raised] very valuable.\n\nAfter our meeting I thought further about [point raised] and [one short idea or answer].\n\nI look forward to hearing from you.\n\nKind regards,\n[Full Name]"),
]
for i, (code, name, use, subj, bodytxt) in enumerate(TPL):
    r = 4+i
    body(wt, r, 1, code, bold=True, color=NAVY); body(wt, r, 2, name, align="left", wrap=True, bold=True)
    body(wt, r, 3, use, align="left", wrap=True); body(wt, r, 4, subj, align="left", wrap=True)
    body(wt, r, 5, bodytxt, align="left", wrap=True, size=9); wt.row_dimensions[r].height = 175
wt.merge_cells("A10:E10"); tc = wt["A10"]; tc.value = "Personalize every email. Always cite one specific paper by the professor, and never send the same text to two people in the same group."
tc.font = Font(name=FONT, size=10, bold=True, color=AMBER); tc.alignment = Alignment(horizontal="center", vertical="center"); wt.row_dimensions[10].height = 28

# =====================================================================
# HOW IT WORKS
# =====================================================================
wg = newsheet(SH_GUIDE, SLATE)
banner(wg, 3, "How It Works", "Built by ApplyAbroadLab.com  |  Free to use and share with attribution")
widths(wg, {"A":26,"B":100,"C":3}); hdr(wg, 3, ["Step","What to do"])
STEPS = [
 ("1. Add positions","Open Position Database and add your target positions. Each ID must be unique; the same ID connects the position to the other sheets. Days to Deadline is calculated and colored automatically."),
 ("2. Log each email","In Communication Tracker, pick the position ID from the drop-down; title, university and country fill in by themselves. Enter the date of your first email and which template and CV version you used."),
 ("3. Check alerts daily","Open Follow-up Alerts every morning. A red row means send a follow-up today, amber means prepare one, green means wait."),
 ("4. Record replies","As soon as a reply arrives, log it in Response Management. Update the Status in Communication Tracker too so the day counter stops."),
 ("5. Watch the dashboard","Seven metrics, four charts, the deadline list and the follow-up list all update on their own. Every card and row is a link: click it to jump to the exact source cell, which is highlighted in a matching color."),
 ("Color guide","Red = urgent or rejected   |   Amber = warning   |   Green = on track   |   Grey cells = formulas, leave them alone   |   Colored column headers = dashboard link targets"),
 ("About the sample data","The 20 positions are real, taken from the ApplyAbroadLab.com database on 27 September 2026. The email dates, statuses and replies are illustrative only so that the charts have something to show. Clear them before real use."),
 ("Full position database","Browse all current PhD and postdoc positions at applyabroadlab.com/position/"),
]
for i, (t, x) in enumerate(STEPS):
    r = 4+i; body(wg, r, 1, t, bold=True, color=NAVY, wrap=True); body(wg, r, 2, x, align="left", wrap=True); wg.row_dimensions[r].height = 44
wg["B11"].hyperlink = "https://applyabroadlab.com/position/"; wg["B11"].font = Font(name=FONT, size=10, color=LINK, underline="single")

wb._sheets = [wb[s] for s in [SH_DASH, SH_POS, SH_COM, SH_RES, SH_AL, SH_TPL, SH_GUIDE]]
wb.properties.creator = "ApplyAbroadLab.com"; wb.properties.lastModifiedBy = "ApplyAbroadLab.com"; wb.properties.title = "PhD & Postdoc Application Tracker"
wb.active = 0
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "phd-postdoc-position-tracker.xlsx")
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
# charts must plot the hidden source columns
import zipfile, shutil
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.startswith("xl/charts/chart"):
            data = data.replace(b'plotVisOnly val="1"', b'plotVisOnly val="0"')
        zout.writestr(item, data)
shutil.move(tmp, OUT); print("saved:", OUT)
